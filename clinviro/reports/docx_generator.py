# ClinViro
# Copyright (C) 2017 Stanford HIVDB team.
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.

# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.

# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <http://www.gnu.org/licenses/>.

# -*- coding: utf-8 -*-
import os
import uuid
from io import BytesIO
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from flask import render_template
from texttable import Texttable
from subprocess import Popen, PIPE
from flask import current_app as app

from .abstract import AbstractGenerator

LINE_WIDTH = 200


class DOCXGenerator(AbstractGenerator):

    content_type = 'docx'
    mimetype = ('application/vnd.openxmlformats-'
                'officedocument.wordprocessingml.document')

    def table_mutation_types(self, genedr):
        table = Texttable(LINE_WIDTH)
        for muttype in genedr['mutation_types']:
            # len(cols) == 2
            row = [muttype['label']]
            muts = [mut['text'].replace('*', '\*').replace('_', '\_')
                    for mut in genedr['mutations']
                    if mut['type'] == muttype['name']]
            row.append(', '.join(muts) or 'None')
            table.add_row(row)
        return table.draw()

    def table_drug_levels(self, drug_levels):
        table = Texttable(LINE_WIDTH)
        for dlevel in drug_levels:
            # len(cols) == 2
            row = [
                '{fullname} ({name})'.format(**dlevel['drug']),
                dlevel['level_text']
            ]
            table.add_row(row)
        return table.draw()

    def table_mutation_scores(self, dcresult):
        dcname = dcresult['drug_class']['name']
        drug_levels = dcresult['drug_levels']
        table = Texttable()
        table.header([dcname] + [dlevel['drug']['name']
                                 for dlevel in drug_levels])
        total = {}
        for mutscore in dcresult['mutation_scores']:
            # len(cols) == len(drugs) + 1
            row = []
            row.append(' + '.join(mutscore['mutations']))
            for dlevel in drug_levels:
                dname = dlevel['drug']['name']
                score = mutscore['drug_scores'][dname]
                row.append(score)
                total[dname] = total.setdefault(dname, 0) + score
            table.add_row(row)
        if len(dcresult['mutation_scores']) != 1:
            row = []
            row.append('Total')
            for dlevel in drug_levels:
                row.append(total.get(dlevel['drug']['name'], 0))
            table.add_row(row)
        return table.draw()

    def table_prev_sequences(self, prevseqs):
        table = Texttable()
        table.header(['Date', 'Accession #', 'Distance %', 'Mutations'])
        for seq in prevseqs:
            # len(cols) == 4
            table.add_row([
                seq['collected_at'],
                seq['vnum'],
                '%0.2f' % (seq['distance'] * 100),
                seq['mutations']
            ])
        return table.draw()

    @staticmethod
    def fix_styles(input_docx):
        doc = Document(input_docx)
        mut_table_style = doc.styles['Mutation Score Table']
        kv_table_style = doc.styles['Key Value Table']
        for table in doc.tables:
            if len(table.row_cells(0)) == 2:
                table.style = kv_table_style
                width = max(len(c.text) for c in table.column_cells(0))
                for cell in table.column_cells(0):
                    cell.width = width * Pt(8)
                width2 = max(len(c.text) for c in table.column_cells(1))
                for cell in table.column_cells(1):
                    cell.width = min(width2, (60 - width)) * Pt(8)
            else:
                table.style = mut_table_style
                if table.cell(0, 1).text == 'Accession #':
                    for i in range(3):
                        width = max(len(c.text) for c in table.column_cells(i))
                        for cell in table.column_cells(i):
                            cell.width = width * Pt(8.5)
        out = BytesIO()
        doc.save(out)
        return out.getvalue()

    def render(self, prepared_data):
        prepared_data = deepcopy(prepared_data)

        mdtext = render_template(
            'docx_report.markdown.jinja2',
            table_mutation_types=self.table_mutation_types,
            table_drug_levels=self.table_drug_levels,
            table_mutation_scores=self.table_mutation_scores,
            table_prev_sequences=self.table_prev_sequences,
            **prepared_data).encode('utf-8')
        stdoutdocx = '/tmp/clin-{}.docx'.format(uuid.uuid4())
        try:
            os.symlink('/dev/stdout', stdoutdocx)
            p = Popen([app.config['CMD_PANDOC'],
                       '-f', 'markdown+grid_tables',
                       '-o', stdoutdocx, '--reference-docx',
                       app.config['PANDOC_REFERENCE_DOCX']],
                      stdout=PIPE, stdin=PIPE, stderr=PIPE)
            stdout, _ = p.communicate(mdtext)
            return self.fix_styles(BytesIO(stdout))
        finally:
            try:
                os.remove(stdoutdocx)
            except FileNotFoundError:
                pass
